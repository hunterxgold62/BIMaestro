"""Read-only Home feasibility probe; never writes the source workbook or models."""
import argparse, collections, datetime, hashlib, json, ntpath, os, re, time
from pathlib import Path
import openpyxl

def extract(source):
    start = time.perf_counter()
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    entries, invalid, rows, sheets, events, kinds = {}, 0, 0, [], collections.Counter(), collections.Counter()
    for sheet in book:
        sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
        if sheet.title != 'Historique_Temps_Revit': continue
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(v or '') for v in next(iterator)]
        def value(row, name):
            try: return row[headers.index(name)]
            except (ValueError, IndexError): return None
        for row in iterator:
            if not any(v is not None for v in row): continue
            rows += 1
            events[str(value(row,'Event') or '')] += 1
            kinds[str(value(row,'Document Kind') or '')] += 1
            candidates = [value(row,n) for n in ['Central Path','Local Path','Document ID']]
            paths=[]
            for raw in candidates:
                path = str(raw or '').split('|')[-1].strip()
                if re.search(r'\.(rvt|rfa)$',path,re.I) and (ntpath.isabs(path) or re.match(r'^(https?://|BIM 360://|Autodesk Docs://|RSN://)',path,re.I)):
                    paths.append(path)
            if not paths:
                invalid += 1
                continue
            path = paths[0]
            key = path.replace('/','\\').casefold()
            date = value(row,'Date')
            if isinstance(date,(datetime.datetime,datetime.date)): date = date.isoformat()[:10]
            when = str(date or '')+'T'+str(value(row,'Time') or '')
            version = str(value(row,'Revit Version') or '')
            item=entries.setdefault(key, {'path':path,'filename':ntpath.basename(path),'extension':ntpath.splitext(path)[1].lower(),'lastLoggedActivity':when,'lastOpened':None,'versions':[],'sourceRows':0,'alternatePaths':[]})
            item['sourceRows'] += 1
            if version and version not in item['versions']: item['versions'].append(version)
            if when > item['lastLoggedActivity']: item['lastLoggedActivity']=when
            if str(value(row,'Event') or '').casefold() in ('opened','open','ouverture','ouvert'):
                if not item['lastOpened'] or when > item['lastOpened']: item['lastOpened']=when
            for alternative in paths[1:]:
                if alternative != path and alternative not in item['alternatePaths']: item['alternatePaths'].append(alternative)
    book.close()
    items = sorted(entries.values(),key=lambda e:e['lastLoggedActivity'],reverse=True)
    return {'entries':items,'summary':{'rows':rows,'unique':len(items),'invalidOrUnresolvableRows':invalid,'extensions':dict(collections.Counter(e['extension'] for e in items)),'unknownVersionEntries':sum(not e['versions'] for e in items),'events':dict(events),'documentKinds':dict(kinds),'sheets':sheets},'elapsedMs':round((time.perf_counter()-start)*1000,3)}

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('source'); parser.add_argument('output'); args=parser.parse_args()
    out=Path(args.output); out.mkdir(parents=True,exist_ok=True)
    source=Path(args.source); before=hashlib.sha256(source.read_bytes()).hexdigest()
    runs=[extract(source) for _ in range(3)]; result=runs[-1]
    cache=source.parent/'CacheVignettes'; cache_start=time.perf_counter(); matched=0
    for item in result['entries']:
        item['thumbnailPath']=None
        for p in [item['path']]+item['alternatePaths']:
            for size in [64,96,128,160,200,256,300,512]:
                candidate=cache/(hashlib.sha1(f'{p}|{size}'.encode()).hexdigest()+'.png')
                if candidate.is_file(): item['thumbnailPath']=str(candidate); break
            if item['thumbnailPath']: break
        matched += bool(item['thumbnailPath'])
    cache_elapsed=round((time.perf_counter()-cache_start)*1000,3)
    checks=[]
    for item in [e for e in result['entries'] if re.match(r'^C:[\\/]',e['path'],re.I)][:12]:
        start=time.perf_counter(); exists=os.path.isfile(item['path']); checks.append({'path':item['path'],'exists':exists,'ms':round((time.perf_counter()-start)*1000,3)})
        item['sampleLocalExists']=exists
    after=hashlib.sha256(source.read_bytes()).hexdigest()
    payload={'source':str(source),'sourceUnchanged':before==after,'generatedAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'summary':result['summary'],'benchmarks':{'workbookReadAndDedupeMs':[r['elapsedMs'] for r in runs],'thumbnailCacheLookupMs':cache_elapsed,'thumbnailCacheMatches':matched,'sampleLocalChecks':checks},'entries':result['entries'],'warnings':['Rows without absolute RVT/RFA path cannot become reopenable catalog items.','lastLoggedActivity is not lastOpened: time logger persists closed activity slices.','Versions record the Revit host used then, not verified current file format.','No model opened or saved; network/cloud existence not probed.','No guarantee of history before BIMaestro logging or after log deletion.']}
    (out/'catalog.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'summary':payload['summary'],'readMs':payload['benchmarks']['workbookReadAndDedupeMs'],'cacheMatches':matched,'cacheMs':cache_elapsed,'localSampleCount':len(checks),'localSampleExists':sum(c['exists'] for c in checks),'sourceUnchanged':payload['sourceUnchanged']},ensure_ascii=False))

if __name__=='__main__': main()
